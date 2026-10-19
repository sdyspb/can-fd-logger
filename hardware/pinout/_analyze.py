import openpyxl

path = r'D:\work\!myprojects\canfd-logger-fc7300\hardware\pinout\FC7300F8MDT Pinout_V1.3.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
print('SHEETS:', wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print('==== SHEET:', name, 'dims:', ws.dimensions, 'rows:', ws.max_row, 'cols:', ws.max_column)
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(25, ws.max_row), values_only=True), 1):
        print(i, [str(c) if c is not None else '' for c in row])
